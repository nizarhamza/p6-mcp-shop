"""Oracles of export-dashboard: P6's PROJECT, PROJWBS, TASK, TASKPRED, TASKRSRC, RSRC and CALENDAR rows read
with SQL, working days counted on a calendar parsed here, the DCMA 14 checks counted again. The card values
are also read back from the written workbook (the tool's output file) and must agree with the SQL values.

Definitions (the dashboard's, as documented): open = not completed; real = not LOE or WBS summary; baseline
dates of an activity = the planned dates of the baseline activity with the same Activity ID; slip = working
days of the activity's calendar from its baseline finish to its finish (actual, else early); Delayed = slip
above near_critical_days (10), At risk = above 0; baseline planned units at the data date = the baseline's
labor and nonlabor assignments spread linearly over calendar time between their planned dates; earned =
budget units x activity % complete (physical %, duration % or units % by the percent complete type).
The DCMA critical path test is not re-run with a CPM engine: it passes here when the earliest most critical
open task reaches the latest finishing activity through successors that are as critical (a delay of that
task then moves the project finish as much)."""

from __future__ import annotations

import math
import os
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal

from oracles import Contains

SUMMARY = ("TT_LOE", "TT_WBS")
MILESTONES = ("TT_Mile", "TT_FinMile")
NEAR = 10.0  # near_critical_days default of p6_export_dashboard
LABOR = ("RT_Labor", "RT_Equip")
HARD = ("CS_MSO", "CS_MEO", "CS_MSOB", "CS_MEOB", "CS_MANDSTART", "CS_MANDFIN")
LIMIT_D = 44.0  # DCMA limit of high float and high duration (days of each calendar)


def _dt(v) -> datetime | None:
    if not v:
        return None
    v = str(v).replace("T", " ")
    return datetime.strptime(v[:19], "%Y-%m-%d %H:%M:%S") if len(v) >= 19 else datetime.strptime(v[:16], "%Y-%m-%d %H:%M")



def _alive_sql(o, table: str, alias: str = "") -> str:
    """' AND <alias>delete_session_id IS NULL' when the table has that column (P6 soft deletes)."""
    cols = {str(c["name"]).lower() for c in o.sql(f"PRAGMA table_info({table})")}
    return f" AND {alias}delete_session_id IS NULL" if "delete_session_id" in cols else ""


# ----------------------------------------------------------------- calendars

def _tree(text: str):
    """P6 text tree (0||name(attrs)(children)) -> (name, [attrs], [children])."""
    pos = 0

    def skip():  # P6 writes line breaks and indents between nodes
        nonlocal pos
        while text[pos] in " \t\r\n":
            pos += 1

    def expect(ch):
        nonlocal pos
        skip()
        if text[pos] != ch:
            raise ValueError(f"calendar text: {ch!r} expected at {pos}")
        pos += 1

    def node():
        nonlocal pos
        expect("(")
        pos = text.index("||", pos) + 2
        op = text.index("(", pos)
        name = text[pos:op].strip()
        cl = text.index(")", op)
        attrs = [a.strip() for a in text[op + 1:cl].split("|")] if text[op + 1:cl].strip() else []
        pos = cl + 1
        expect("(")
        kids = []
        skip()
        while text[pos] == "(":
            kids.append(node())
            skip()
        expect(")")
        expect(")")
        return name, attrs, kids
    return node()


def _minutes(hhmm: str) -> int:
    h, m = hhmm.split(":")
    return int(h) * 60 + int(m)


def _shift(attrs):
    kv = dict(zip(attrs[::2], attrs[1::2], strict=False))
    s, f = _minutes(kv["s"]), _minutes(kv["f"])
    return (s, f or 1440) if s <= (f or 1440) else (f, s)


class Cal:
    """Working time of a calendar: week (Python weekday -> shifts in minutes) and exceptions by date."""

    def __init__(self, o, clndr_id, depth=0):
        """A missing or unreadable calendar has error set (no working time, nothing from its base)."""
        rows = o.sql("SELECT clndr_data, day_hr_cnt, base_clndr_id FROM CALENDAR WHERE clndr_id = ?"
                     + _alive_sql(o, "CALENDAR"), clndr_id) if clndr_id is not None else []
        self.week, self.exc, self.error, self.hpd = {k: () for k in range(7)}, {}, None, 8.0
        if not rows:
            self.error = "missing"
            return
        r = rows[0]
        self.hpd = float(r["day_hr_cnt"] or 0) or 8.0
        data = r["clndr_data"]
        try:
            tree = _tree(data.decode("utf-8", "replace") if isinstance(data, bytes) else data)
            week, exc = None, {}
            for name, _a, kids in tree[2]:
                if name == "DaysOfWeek":
                    week = {k: () for k in range(7)}
                    for day, _da, shifts in kids:
                        week[(int(day) + 5) % 7] = tuple(sorted(_shift(sa) for _n, sa, _k in shifts))
                elif name == "Exceptions":
                    for _n, attrs, shifts in kids:
                        kv = dict(zip(attrs[::2], attrs[1::2], strict=False))
                        d = date(1899, 12, 30) + timedelta(days=int(float(kv["d"])))
                        exc[d] = tuple(sorted(_shift(sa) for _n2, sa, _k in shifts))
        except (ValueError, KeyError, IndexError, TypeError, AttributeError) as e:
            self.error = f"unreadable: {e}"
            return
        if r["base_clndr_id"] is not None and depth < 10:
            base = Cal(o, r["base_clndr_id"], depth + 1)
            if not base.error:
                self.week, self.exc = dict(base.week), dict(base.exc)
        if week is not None:
            self.week = week
        self.exc.update(exc)  # the calendar's own dates win

    def shifts(self, d: date):
        return self.exc[d] if d in self.exc else self.week[d.weekday()]

    def work_hours(self, a: datetime, b: datetime) -> float:
        if b < a:
            return -self.work_hours(b, a)
        secs, d = 0.0, a.date()
        while d <= b.date():
            day0 = datetime.combine(d, datetime.min.time())
            for s, f in self.shifts(d):
                lo, hi = max(a, day0 + timedelta(minutes=s)), min(b, day0 + timedelta(minutes=f))
                if hi > lo:
                    secs += (hi - lo).total_seconds()
            d += timedelta(days=1)
        return secs / 3600

    def wd(self, a, b):
        """Working days from a to b, 1 decimal (None without both dates); calendar days when the calendar is
        missing or unreadable (the dashboard does the same and warns)."""
        if a is None or b is None:
            return None
        if self.error:
            return round((b - a).total_seconds() / 86400, 1) + 0.0
        return round(self.work_hours(a, b) / self.hpd, 1) + 0.0


_CAL: dict = {}


def cal(o, clndr_id) -> Cal:
    key = (o.db, clndr_id)
    if key not in _CAL:
        _CAL[key] = Cal(o, clndr_id)
    return _CAL[key]


# ----------------------------------------------------------------- rows

def _project(o, pid=None):
    return o.sql("SELECT * FROM PROJECT WHERE proj_id = ?", pid or o.pid)[0]


def _bid(o):
    if o.args.get("baseline"):
        return o.one("SELECT proj_id FROM PROJECT WHERE proj_short_name = ?", o.args["baseline"])
    return _project(o)["sum_base_proj_id"]


def _tasks(o, pid):
    return o.sql("SELECT t.*, COALESCE(NULLIF(c.day_hr_cnt, 0), 8) AS hpd FROM TASK t LEFT JOIN CALENDAR c ON "
                 "c.clndr_id = t.clndr_id WHERE t.proj_id = ? AND t.delete_session_id IS NULL ORDER BY t.task_id", pid)


def _units_asg(o, pid):
    """Labor and nonlabor assignments of a project with their activity."""
    return o.sql("SELECT tr.target_qty, tr.target_start_date, tr.target_end_date, t.phys_complete_pct, t.status_code, "
                 "t.complete_pct_type, t.target_drtn_hr_cnt, t.remain_drtn_hr_cnt, "
                 "COALESCE(t.act_work_qty, 0) + COALESCE(t.act_equip_qty, 0) AS t_act, "
                 "COALESCE(t.remain_work_qty, 0) + COALESCE(t.remain_equip_qty, 0) AS t_rem "
                 "FROM TASKRSRC tr JOIN TASK t ON t.task_id = tr.task_id LEFT JOIN RSRC r ON r.rsrc_id = tr.rsrc_id "
                 "WHERE t.proj_id = ? AND tr.delete_session_id IS NULL AND t.delete_session_id IS NULL "
                 "AND COALESCE(tr.rsrc_type, r.rsrc_type) IN ('RT_Labor', 'RT_Equip')", pid)


def _perf(r) -> float:
    done = 100.0 if r["status_code"] == "TK_Complete" else 0.0
    if r["complete_pct_type"] == "CP_Drtn":
        od = float(r["target_drtn_hr_cnt"] or 0)
        v = 100 * (od - float(r["remain_drtn_hr_cnt"] or 0)) / od if od else done
    elif r["complete_pct_type"] == "CP_Units":
        at = float(r["t_act"]) + float(r["t_rem"])
        v = 100 * float(r["t_act"]) / at if at else done
    else:
        v = float(r["phys_complete_pct"] or 0)
    return min(100.0, max(0.0, v))


def _planned_at(rows, t: datetime) -> tuple[float, float]:
    """(planned quantity at t, budget) of assignments spread linearly between their planned dates."""
    got = budget = 0.0
    for r in rows:
        q = float(r["target_qty"] or 0)
        budget += q
        s, e = _dt(r["target_start_date"]), _dt(r["target_end_date"])
        if not (q and s and e):
            continue
        e = max(s, e)
        if t >= e and t > s:
            got += q
        elif s < t < e:
            got += q * (t - s).total_seconds() / (e - s).total_seconds()
    return got, budget


# ----------------------------------------------------------------- cards

def _figures(o) -> dict:
    if getattr(o, "_dash_figures", None) is None:
        o._dash_figures = _compute(o)
    return o._dash_figures


def _compute(o) -> dict:
    p, bid = _project(o), _bid(o)
    dd, scd = _dt(p["last_recalc_date"]), _dt(p["scd_end_date"])
    pcal = cal(o, p["clndr_id"])
    acts = _tasks(o, o.pid)
    act = sum(float(a["act_work_qty"] or 0) + float(a["act_equip_qty"] or 0) for a in acts)
    rem = sum(float(a["remain_work_qty"] or 0) + float(a["remain_equip_qty"] or 0) for a in acts)
    units = 100 * act / (act + rem) if act + rem else None
    f = {"units": units, "dd": dd, "scd": scd, "pcal": pcal, "bid": bid, "acts": acts}
    thr = float(p["critical_drtn_hr_cnt"] or 0)
    f["thr_d"] = thr / pcal.hpd
    crit = (lambda a: a["driving_path_flag"] == "Y") if p["critical_path_type"] == "CT_DrivPath" else \
        (lambda a: a["total_float_hr_cnt"] is not None and a["total_float_hr_cnt"] <= thr)
    f["critical"] = sum(1 for a in acts if a["status_code"] != "TK_Complete" and a["task_type"] not in SUMMARY
                        and crit(a))
    f["remaining"] = max(0.0, pcal.wd(dd, scd)) if dd and scd else None
    found = {}
    for a in sorted((a for a in acts if a["task_type"] in MILESTONES), key=lambda a: (a["task_code"], a["task_id"])):
        found.setdefault(a["task_code"], a)

    def mdate(m, typ, planned=False):
        kind = "start" if typ == "TT_Mile" else "end"
        return _dt(m[f"target_{kind}_date"] if planned else m[f"act_{kind}_date"] or m[f"early_{kind}_date"])
    ms = sorted(found.values(), key=lambda m: (mdate(m, m["task_type"]) or datetime.max, m["task_code"]))[:30]
    f["milestones"] = len(ms)
    if bid is None:
        f.update(plan=None, delayed=None, at_risk=None, ev=None, bp=None, bl_finish=None, vs_bl=None, ms_open=0,
                 ms_late=None)
        return f
    base = {}
    for b in sorted(_tasks(o, bid), key=lambda b: -b["task_id"]):  # lowest task_id wins
        base[b["task_code"]] = b
    bp, bb = _planned_at(_units_asg(o, bid), dd)
    f["bp"], f["plan"] = bp, round(100 * bp / bb, 4) if bb else None
    f["ev"] = sum(float(r["target_qty"] or 0) * _perf(r) / 100 for r in _units_asg(o, o.pid))
    slips = []
    for a in acts:
        if a["status_code"] == "TK_Complete" or a["task_type"] in SUMMARY:
            continue
        b = base.get(a["task_code"])
        bf, fin = _dt(b["target_end_date"]) if b else None, _dt(a["act_end_date"] or a["early_end_date"])
        slips.append(cal(o, a["clndr_id"]).wd(bf, fin))
    f["delayed"] = sum(1 for s in slips if s is not None and s > NEAR)
    f["at_risk"] = sum(1 for s in slips if s is not None and 0 < s <= NEAR)
    f["bl_finish"] = _dt(o.one("SELECT MAX(target_end_date) FROM TASK WHERE proj_id = ? AND delete_session_id IS NULL",
                               bid))
    f["vs_bl"] = pcal.wd(f["bl_finish"], scd)
    var = []
    for m in ms:
        b = base.get(m["task_code"])
        if m["status_code"] == "TK_Complete" or b is None:
            continue
        v = cal(o, m["clndr_id"]).wd(mdate(b, m["task_type"], True), mdate(m, m["task_type"]))
        if v is not None:
            var.append(v)
    f["ms_open"], f["ms_late"] = len(var), sum(1 for v in var if v > 0) if var else None
    return f


def _workbook_cards(o) -> list[tuple]:
    """(value, detail) of each card as written in rows 6 and 8 of the Dashboard sheet."""
    import openpyxl
    ws = openpyxl.load_workbook(o.args["path"])["Dashboard"]
    grid = range(3, 30)  # columns C to AC: the cards (helper data of the charts lies further right)
    vals = [ws.cell(6, c).value for c in grid if ws.cell(6, c).value is not None]
    details = [ws.cell(8, c).value for c in grid if ws.cell(8, c).value is not None]
    return list(zip(vals, details, strict=True))


def cards(o):
    f = _figures(o)
    has_bl = f["bid"] is not None
    sv = spi = None
    if has_bl and f["bp"]:
        bp, var = round(f["bp"], 2), round(f["ev"] - f["bp"], 2)
        sv, spi = 100 * var / bp, round(f["ev"] / f["bp"], 4)
    plan, u = f["plan"], f["units"]
    rows = [
        ("Overall Progress", u, Contains(f"Planned: {plan:.2f}%", f"Variance: {u - plan:+.2f}%")
         if plan is not None and u is not None else Contains("No baseline plan")),
        ("Critical Activities", f["critical"], Contains(f"Total float <= {f['thr_d']:g} d")),
        ("Delayed Activities", f["delayed"], Contains(f"over {NEAR:g} wd", f"{f['at_risk']:,} at risk")
         if has_bl else Contains("No baseline")),
        ("Schedule Variance", sv, Contains(f"SPI {spi:.2f}") if sv is not None else Contains("baseline")),
        ("Forecast Completion", f["scd"].strftime("%Y-%m-%d"),
         Contains(f"Baseline: {f['bl_finish']:%d %b %Y}", f"({f['vs_bl']:+.1f} wd)") if has_bl
         else Contains("No baseline")),
        ("Remaining Duration", f["remaining"], Contains("Working days to the forecast finish")),
        ("Milestones Late", f["ms_late"], Contains(f"of {f['ms_open']} open milestones with a baseline")
         if f["ms_open"] else Contains("No open milestone"))]
    # the workbook (the tool's output file) must hold the same values and details
    book = _workbook_cards(o)
    assert len(book) == 7, f"{len(book)} card values in the Dashboard sheet"
    for (name, value, detail), (bv, bd) in zip(rows, book, strict=True):
        if isinstance(value, float):
            assert bv != "n/a" and abs(float(bv) - value) < 0.0051, f"workbook card {name}: {bv!r}, SQL {value!r}"
        elif isinstance(bv, datetime):
            assert bv.strftime("%Y-%m-%d") == value, f"workbook card {name}: {bv!r}, SQL {value!r}"
        else:
            assert bv == ("n/a" if value is None else value), f"workbook card {name}: {bv!r}, SQL {value!r}"
        assert detail(bd), f"workbook card {name} detail {bd!r} does not satisfy {detail!r}"
    return [{"card": n, "value": _shown(n, v), "detail": d} for n, v, d in rows]


def _half_up(v: float, nd: int) -> str:
    """abs(v) rounded half away from zero to nd decimals, as Excel displays it."""
    return f"{Decimal(str(abs(v))).quantize(Decimal(1).scaleb(-nd), rounding=ROUND_HALF_UP):,.{nd}f}"


def _shown(card: str, v) -> str:
    """The card value as Excel shows it with the card's number format ('n/a' when there is none)."""
    if v is None:
        return "n/a"
    if card == "Overall Progress":  # 0.00"%"
        return ("-" if v < 0 else "") + _half_up(v, 2).replace(",", "") + "%"
    if card == "Schedule Variance":  # +0.0"%";-0.0"%";0.0"%"
        return ("+" if v > 0 else "-" if v < 0 else "") + _half_up(v, 1).replace(",", "") + "%"
    if card == "Forecast Completion":  # dd-mmm-yy
        return datetime.strptime(v, "%Y-%m-%d").strftime("%d-%b-%y")
    if card == "Remaining Duration":  # 0" d" (0.0" d" when not whole)
        return _half_up(v, 0 if float(v).is_integer() else 1).replace(",", "") + " d"
    return f"{v:,}"  # #,##0 and 0: whole counts


# ----------------------------------------------------------------- DCMA 14

# DCMA 14 check -> (DCMA default threshold, rule, PROJPROP 'schedule_check' target key, limit key (hours), on/off key)
DCMA = {
    "logic": (5, "max", "logic_target", None, "check_logic"),
    "leads": (0, "max", "negative_lag_target", None, "check_negative_lag"),
    "lags": (5, "max", "lag_target", None, "check_lag"),
    "long_lags": (None, "max", "long_lags_target", "long_lags_constraint", "check_long_lags"),
    "fs_links": (90, "min", "relationship_target", None, "check_relationships"),
    "hard_constraints": (5, "max", "hard_constraint_target", None, "check_hard_constraints"),
    "high_float": (5, "max", "large_float_target", "large_float_constraint", "check_large_float"),
    "negative_float": (0, "max", "negative_float_target", None, "check_negative_float"),
    "high_duration": (5, "max", "large_duration_target", "large_durations_constraint", "check_large_duration"),
    "invalid_dates": (0, "max", "progress_date_target", None, "check_progress_duration"),
    "resources": (None, "max", "resouces_target", None, "check_resources"),  # P6's spelling
    "missed_tasks": (5, "max", "late_activities_target", None, "check_late_activities"),
    "critical_path_test": (None, None, "cpt_target", None, "check_cpt"),
    "cpli": (0.95, "ratio", "cpli_target", None, "check_cpli"),
    "bei": (0.95, "ratio", "bei_target", None, "check_bei"),
}


def _num(v):
    try:
        f = float(v)
    except (TypeError, ValueError):
        return None
    return f if math.isfinite(f) else None


def _props(o, name: str) -> dict:
    """Top attributes of a PROJPROP text tree of the project (keys lower case); {} when absent or unreadable."""
    v = o.one(f"SELECT prop_value FROM PROJPROP WHERE proj_id = ? AND prop_name = ?{_alive_sql(o, 'PROJPROP')}",
              o.pid, name)
    if v is None:
        return {}
    try:
        attrs = _tree(v.decode("utf-8", "replace") if isinstance(v, bytes) else v)[1]
    except (ValueError, IndexError):
        return {}
    return {str(k).lower(): x for k, x in zip(attrs[::2], attrs[1::2], strict=False)}


def _settings(o) -> tuple[dict, dict, set]:
    """(threshold per check, limit per check as ('h', hours) or ('d', days), disabled checks): the project's
    Check Schedule settings when stored (cpli_target is a tolerance in %), else the DCMA defaults."""
    p = _props(o, "schedule_check")
    thr, lim, off = {}, {}, set()
    for name, (default, _rule, key, lkey, flag) in DCMA.items():
        t = default if _num(p.get(key)) is None else _num(p[key])
        if name == "cpli" and _num(p.get(key)) is not None:
            t = round(1 - t / 100, 4)
        if p.get(flag) == "N":
            t = None
            off.add(name)
        thr[name] = t
        if lkey:
            lim[name] = ("h", _num(p[lkey])) if _num(p.get(lkey)) is not None else ("d", LIMIT_D)
    return thr, lim, off


def _dcma_failed(o) -> list[str]:
    thr, lim, off = _settings(o)
    p, bid = _project(o), _bid(o)
    dd = _dt(p["last_recalc_date"])
    allt = _tasks(o, o.pid)
    t_by = {a["task_id"]: a for a in allt}
    acts = [a for a in allt if a["task_type"] not in SUMMARY]
    opened = [a for a in acts if a["status_code"] != "TK_Complete"]
    work = [a for a in opened if a["task_type"] in ("TT_Task", "TT_Rsrc")]
    alive = _alive_sql(o, "TASKPRED", "r.") + _alive_sql(o, "TASK", "t.") + _alive_sql(o, "TASK", "p.")
    links = o.sql("SELECT r.task_id, r.pred_task_id, r.pred_type, r.lag_hr_cnt, p.task_type AS ptype, "
                  "COALESCE(NULLIF(cp.day_hr_cnt, 0), 8) AS hp, COALESCE(NULLIF(cs.day_hr_cnt, 0), 8) AS hs "
                  "FROM TASKPRED r JOIN TASK t ON t.task_id = r.task_id JOIN TASK p ON p.task_id = r.pred_task_id "
                  "LEFT JOIN CALENDAR cp ON cp.clndr_id = p.clndr_id LEFT JOIN CALENDAR cs ON cs.clndr_id = t.clndr_id "
                  "WHERE t.proj_id = ?" + alive, o.pid)
    olinks = [lk for lk in links if t_by[lk["task_id"]]["status_code"] != "TK_Complete"
              and t_by[lk["task_id"]]["task_type"] not in SUMMARY and lk["ptype"] not in SUMMARY]
    alive = _alive_sql(o, "TASKPRED", "r.") + _alive_sql(o, "TASK", "s.") + _alive_sql(o, "TASK", "p.")
    real_links = o.sql("SELECT r.task_id, r.pred_task_id FROM TASKPRED r JOIN TASK s ON s.task_id = r.task_id "
                       "JOIN TASK p ON p.task_id = r.pred_task_id WHERE s.task_type NOT IN ('TT_LOE', 'TT_WBS') "
                       "AND p.task_type NOT IN ('TT_LOE', 'TT_WBS')" + alive)
    has_pred = {r["task_id"] for r in real_links}
    has_succ = {r["pred_task_id"] for r in real_links}
    starts = sorted((a for a in allt if a["task_type"] == "TT_Mile" and a["task_id"] not in has_pred),
                    key=lambda a: (a["act_start_date"] or a["early_start_date"] or "9999", a["task_code"], a["task_id"]))
    fins = sorted((a for a in allt if a["task_type"] == "TT_FinMile" and a["task_id"] not in has_succ),
                  key=lambda a: (a["task_code"], a["task_id"]))
    fins = sorted(fins, key=lambda a: a["act_end_date"] or a["early_end_date"] or "", reverse=True)
    sms, fms = starts[0]["task_id"] if starts else None, fins[0]["task_id"] if fins else None
    res = {}  # check -> pass (None: not evaluated)

    def check(name, n, d, value=None):
        t, rule = thr[name], DCMA[name][1]
        v = value if rule == "ratio" else (100 * n / d if d else None)
        res[name] = None if t is None or v is None else v >= t if rule in ("min", "ratio") else v <= t

    def over(hours, hpd, limit):  # above a limit in hours, or in days of the given hours per day
        return hours is not None and hours > (limit[1] if limit[0] == "h" else limit[1] * hpd)
    nopen, nl = len(opened), len(olinks)
    check("logic", sum(1 for a in opened if (a["task_id"] not in has_pred and a["task_id"] != sms)
                       or (a["task_id"] not in has_succ and a["task_id"] != fms)), nopen)
    check("leads", sum(1 for lk in olinks if (lk["lag_hr_cnt"] or 0) < 0), nl)
    check("lags", sum(1 for lk in olinks if (lk["lag_hr_cnt"] or 0) > 0), nl)
    rule = _props(o, "scheduling").get("sched_calendar_on_relationship_lag") or "rcal_Predecessor"
    phpd = cal(o, p["clndr_id"]).hpd

    def lag_hpd(lk):  # hours per lag day under the project's lag calendar option
        return {"rcal_Successor": lk["hs"], "rcal_24Hour": 24.0, "rcal_ProjDefault": phpd}.get(rule, lk["hp"])
    check("long_lags", sum(1 for lk in olinks if over(lk["lag_hr_cnt"], lag_hpd(lk), lim["long_lags"])), nl)
    check("fs_links", sum(1 for lk in olinks if lk["pred_type"] == "PR_FS"), nl)
    check("hard_constraints", sum(1 for a in opened if a["cstr_type"] in HARD or a["cstr_type2"] in HARD), nopen)
    check("high_float", sum(1 for a in opened if over(a["total_float_hr_cnt"], a["hpd"], lim["high_float"])), nopen)
    check("negative_float", sum(1 for a in opened if a["total_float_hr_cnt"] is not None
                                and a["total_float_hr_cnt"] < 0), nopen)
    check("high_duration", sum(1 for a in work if over(a["remain_drtn_hr_cnt"], a["hpd"], lim["high_duration"])),
          len(work))
    if dd:
        bad = sum(1 for a in acts if any(x and _dt(x) > dd for x in (a["act_start_date"], a["act_end_date"]))
                  or (a["status_code"] != "TK_Complete" and any(x and _dt(x) < dd for x in (a["early_start_date"],
                                                                                             a["early_end_date"]))))
        check("invalid_dates", bad, len(acts))
    else:
        res["invalid_dates"] = None
    asg = {r["task_id"] for r in o.sql("SELECT a.task_id FROM TASKRSRC a JOIN TASK t ON t.task_id = a.task_id "
                                       "WHERE t.proj_id = ?" + _alive_sql(o, "TASKRSRC", "a."), o.pid)}
    check("resources", sum(1 for a in work if (a["target_drtn_hr_cnt"] or 0) > 0 and a["task_id"] not in asg),
          len(work))
    if bid is not None and dd:
        cur = {}
        for a in sorted(allt, key=lambda a: -a["task_id"]):
            cur[a["task_code"]] = a
        due = [b for b in _tasks(o, bid) if b["task_type"] not in SUMMARY and b["target_end_date"]
               and _dt(b["target_end_date"]) <= dd]

        def missed(b):
            a = cur.get(b["task_code"])
            if a is None or a["status_code"] != "TK_Complete":
                return True
            return bool(a["act_end_date"]) and _dt(a["act_end_date"]).date() > _dt(b["target_end_date"]).date()
        check("missed_tasks", sum(1 for b in due if missed(b)), len(due))
        done = sum(1 for a in acts if a["status_code"] == "TK_Complete")
        check("bei", None, None, done / len(due) if due else None)
    else:
        res["missed_tasks"] = res["bei"] = None
    res["critical_path_test"] = None if "critical_path_test" in off else _cpt(o, allt, links)
    pcal = cal(o, p["clndr_id"])
    end, must = _dt(p["scd_end_date"]), _dt(p["plan_end_date"])
    value = None
    if dd and end and must and end > dd:
        cpl = 0 if pcal.error else pcal.work_hours(dd, end)
        if cpl > 0:
            value = (cpl + pcal.work_hours(end, must)) / cpl
        else:  # calendar days without a readable project calendar
            value = ((end - dd) + (must - end)) / (end - dd)
    check("cpli", None, None, value)
    return [k for k in DCMA if res[k] is False]


def _cpt(o, allt, links):
    """The earliest most critical open task (total float within 1 h of the lowest, by start then Activity
    ID) reaches the latest finishing activity through successors as critical: True."""
    expect = _props(o, "scheduling").get("sched_use_expect_end_flag") == "Y"
    cand = [a for a in allt if a["task_type"] in ("TT_Task", "TT_Rsrc") and a["status_code"] != "TK_Complete"
            and not (expect and a["expect_end_date"]) and "CS_MANDFIN" not in (a["cstr_type"], a["cstr_type2"])
            and a["total_float_hr_cnt"] is not None]
    if not cand:
        return None
    low = min(a["total_float_hr_cnt"] for a in cand)
    d = min((a for a in cand if a["total_float_hr_cnt"] <= low + 1),
            key=lambda a: (_dt(a["act_start_date"] or a["early_start_date"]) or datetime.max, a["task_code"]))
    own = [a for a in allt if a["task_type"] not in SUMMARY]
    ends = [x for x in (_dt(a["act_end_date"] or a["early_end_date"]) for a in own) if x]
    if not ends:
        return None
    last = max(ends)
    crit = {a["task_id"] for a in own if a["total_float_hr_cnt"] is not None and a["total_float_hr_cnt"] <= low + 1}
    succ: dict = {}
    for lk in links:
        succ.setdefault(lk["pred_task_id"], []).append(lk["task_id"])
    t_by = {a["task_id"]: a for a in allt}
    seen, stack = {d["task_id"]}, [d["task_id"]]
    while stack:
        n = stack.pop()
        if _dt(t_by[n]["act_end_date"] or t_by[n]["early_end_date"]) == last:
            return True
        for s in succ.get(n, []):
            if s in crit and s not in seen:
                seen.add(s)
                stack.append(s)
    return False


# ----------------------------------------------------------------- warnings, slots

def warnings(o):
    f = _figures(o)
    out = []
    n_wbs = o.one("SELECT COUNT(*) FROM PROJWBS WHERE proj_id = ? AND delete_session_id IS NULL", o.pid)
    if n_wbs > 5000:
        out.append(Contains("More than 5000 WBS nodes"))
    if f["milestones"] == 30 and sum(1 for a in f["acts"] if a["task_type"] in MILESTONES) > 30:
        out.append(Contains("milestones: the first 30"))
    if len(f["acts"]) > 20000:
        out.append(Contains("the Schedule Overview shows the first"))
    if not any(float(r["target_qty"] or 0) for r in _units_asg(o, o.pid)):
        out.append(Contains("S-curve uses cost"))
    return out


def _bl_name(o):
    bid = _bid(o)
    return None if bid is None else o.one("SELECT proj_short_name FROM PROJECT WHERE proj_id = ?", bid)


def _counts(o):
    f = _figures(o)
    return {"wbs": o.one("SELECT COUNT(*) FROM PROJWBS WHERE proj_id = ? AND delete_session_id IS NULL", o.pid),
            "activities": len(f["acts"]), "milestones": f["milestones"], "critical": f["critical"]}


def _file(o):
    path = os.path.abspath(o.args["path"])
    return path if os.path.isfile(path) else f"(no file at {path})"


SLOTS = {
    "data_date": lambda o: _project(o)["last_recalc_date"],
    "last_f9": lambda o: _project(o)["last_schedule_date"],
    "stale": lambda o: o.stale(),
    "bl_used": _bl_name,
    "file": _file,
    "n_acts": lambda o: _counts(o)["activities"],
    "n_wbs": lambda o: _counts(o)["wbs"],
    "n_ms": lambda o: _counts(o)["milestones"],
    "n_crit": lambda o: _counts(o)["critical"],
}
TABLES = {
    "cards": cards,
    "dcma_failed": lambda o: [{"check": c} for c in _dcma_failed(o)],
    "warnings": lambda o: [{"warning": w} for w in warnings(o)],
}
